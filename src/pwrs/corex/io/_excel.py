# Copyright (c) 2026, Liangyu Zhang
# SPDX-License-Identifier: BSD-3-Clause
"""Human-reviewable Excel workbook case format."""

from __future__ import annotations

import math
from collections.abc import Mapping
from os import PathLike
from pathlib import Path
from typing import Any

import numpy as np
from scipy import sparse

from ..mpc import MatpowerCase
from ._schema import (
    DENSE_MATRIX_FIELDS,
    MATRIX_FIELDS,
    SCALAR_FIELDS,
    SCHEMA_NAME,
    SCHEMA_VERSION,
    STRING_LIST_FIELDS,
    VECTOR_FIELDS,
    case_mapping,
    normalize_case,
)

_META_HEADER = ("field", "kind", "dtype", "shape", "value")
_KNOWN_HEADERS = {
    "bus": (
        "bus_i", "type", "Pd", "Qd", "Gs", "Bs", "area", "Vm", "Va",
        "baseKV", "zone", "Vmax", "Vmin", "lam_P", "lam_Q", "mu_Vmax", "mu_Vmin",
    ),
    "gen": (
        "bus", "Pg", "Qg", "Qmax", "Qmin", "Vg", "mBase", "status", "Pmax", "Pmin",
        "Pc1", "Pc2", "Qc1min", "Qc1max", "Qc2min", "Qc2max", "ramp_agc", "ramp_10",
        "ramp_30", "ramp_q", "apf", "mu_Pmax", "mu_Pmin", "mu_Qmax", "mu_Qmin",
    ),
    "branch": (
        "fbus", "tbus", "r", "x", "b", "rateA", "rateB", "rateC", "ratio", "angle",
        "status", "angmin", "angmax", "Pf", "Qf", "Pt", "Qt", "mu_Sf", "mu_St",
        "mu_angmin", "mu_angmax",
    ),
    "areas": ("area", "refbus"),
    "dcline": (
        "fbus", "tbus", "status", "Pf", "Pt", "Qf", "Qt", "Vf", "Vt", "Pmin",
        "Pmax", "QminF", "QmaxF", "QminT", "QmaxT", "loss0", "loss1", "mu_Pmin",
        "mu_Pmax", "mu_QminF", "mu_QmaxF", "mu_QminT", "mu_QmaxT",
    ),
    "fparm": ("dd", "rh", "kk", "mm"),
}


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "Excel case IO requires openpyxl; install pwrs with the 'all' dependency group"
        ) from exc
    return openpyxl


def _path(value: str | PathLike[str]) -> Path:
    path = Path(value)
    if not path.suffix:
        return path.with_suffix(".xlsx")
    if path.suffix.lower() == ".xls":
        raise ValueError("legacy .xls workbooks are not supported; use .xlsx")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"expected a .xlsx path, got {path}")
    return path


def _shape_text(shape: tuple[int, ...]) -> str:
    return ",".join(str(item) for item in shape)


def _parse_shape(field: str, value: Any, dimensions: int) -> tuple[int, ...]:
    try:
        shape = tuple(int(item) for item in str(value).split(",") if item != "")
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Excel field {field!r} has an invalid shape") from exc
    if len(shape) != dimensions or any(item < 0 for item in shape):
        raise ValueError(f"Excel field {field!r} has an invalid shape")
    return shape


def _cell(value: Any) -> Any:
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        if math.isnan(value):
            return "NaN"
        return "Infinity" if value > 0 else "-Infinity"
    return value


def _headers(name: str, columns: int) -> list[str]:
    if name in {"gencost", "dclinecost"}:
        known = ["model", "startup", "shutdown", "ncost"]
        return (known + [f"cost_{index}" for index in range(1, columns - 3)])[:columns]
    known = list(_KNOWN_HEADERS.get(name, ()))
    return (known + [f"col_{index}" for index in range(len(known) + 1, columns + 1)])[:columns]


def _append_dense_sheet(workbook: Any, name: str, value: np.ndarray) -> None:
    sheet = workbook.create_sheet(name)
    headers = _headers(name, value.shape[1])
    sheet.append(headers or ["__empty__"])
    for row in value:
        sheet.append([_cell(item) for item in row])


def _append_vector_sheet(workbook: Any, name: str, value: np.ndarray) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(["value"])
    for item in value:
        sheet.append([_cell(item)])


def _append_sparse_sheet(workbook: Any, name: str, value: Any) -> None:
    matrix = value.tocoo()
    sheet = workbook.create_sheet(name)
    sheet.append(["row", "column", "value"])
    for row, column, item in zip(matrix.row, matrix.col, matrix.data):
        sheet.append([int(row), int(column), _cell(item)])


def write_excel(
    case: MatpowerCase | Mapping[str, Any], destination: str | PathLike[str]
) -> Path:
    """Write one field per sheet with named MATPOWER columns and metadata."""
    openpyxl = _openpyxl()
    path = _path(destination)
    data = case_mapping(case)
    workbook = openpyxl.Workbook()
    meta = workbook.active
    if meta is None:
        raise RuntimeError("openpyxl created a workbook without an active sheet")
    meta.title = "_meta"
    meta.append(_META_HEADER)
    meta.append(["__schema__", "schema", "", "", SCHEMA_NAME])
    meta.append(["__schema_version__", "schema", "", "", SCHEMA_VERSION])

    for name, value in data.items():
        if name in SCALAR_FIELDS:
            meta.append([name, "scalar", type(value).__name__, "", _cell(value)])
        elif name in DENSE_MATRIX_FIELDS:
            array = np.asarray(value)
            meta.append([name, "dense", str(array.dtype), _shape_text(array.shape), ""])
            _append_dense_sheet(workbook, name, array)
        elif name in MATRIX_FIELDS and sparse.issparse(value):
            meta.append([name, "sparse", str(value.dtype), _shape_text(value.shape), ""])
            _append_sparse_sheet(workbook, name, value)
        elif name in MATRIX_FIELDS:
            array = np.asarray(value)
            meta.append([name, "dense", str(array.dtype), _shape_text(array.shape), ""])
            _append_dense_sheet(workbook, name, array)
        elif name in VECTOR_FIELDS:
            array = np.asarray(value).reshape(-1)
            meta.append([name, "vector", str(array.dtype), _shape_text(array.shape), ""])
            _append_vector_sheet(workbook, name, array)
        elif name in STRING_LIST_FIELDS:
            meta.append([name, "strings", "str", str(len(value)), ""])
            sheet = workbook.create_sheet(name)
            sheet.append(["value"])
            for item in value:
                sheet.append([item])
        else:
            raise AssertionError(f"unhandled Excel case field: {name}")

    workbook.save(path)
    return path


def _rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def _dtype(field: str, value: Any) -> np.dtype[Any]:
    try:
        dtype = np.dtype(str(value))
    except TypeError as exc:
        raise ValueError(f"Excel field {field!r} has an invalid dtype") from exc
    if not np.issubdtype(dtype, np.number):
        raise ValueError(f"Excel field {field!r} must have a numeric dtype")
    return dtype


def _read_dense(field: str, sheet: Any, dtype: np.dtype[Any], shape: tuple[int, int]) -> np.ndarray:
    rows = _rows(sheet)
    expected_header = tuple(_headers(field, shape[1]) or ["__empty__"])
    if not rows or rows[0] != expected_header:
        raise ValueError(f"Excel sheet {field!r} has an invalid header")
    values = rows[1:]
    if len(values) != shape[0] or any(len(row) != max(shape[1], 1) for row in values):
        raise ValueError(f"Excel sheet {field!r} does not match declared shape {shape}")
    if shape[1] == 0:
        return np.empty(shape, dtype=dtype)
    try:
        return np.asarray(values, dtype=dtype).reshape(shape)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Excel sheet {field!r} contains non-numeric data") from exc


def _read_vector(field: str, sheet: Any, dtype: np.dtype[Any], shape: tuple[int]) -> np.ndarray:
    rows = _rows(sheet)
    if not rows or rows[0] != ("value",) or len(rows) - 1 != shape[0]:
        raise ValueError(f"Excel sheet {field!r} does not match declared shape {shape}")
    try:
        return np.asarray([row[0] for row in rows[1:]], dtype=dtype)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Excel sheet {field!r} contains non-numeric data") from exc


def _read_sparse(field: str, sheet: Any, dtype: np.dtype[Any], shape: tuple[int, int]) -> Any:
    rows = _rows(sheet)
    if not rows or rows[0] != ("row", "column", "value"):
        raise ValueError(f"Excel sheet {field!r} has an invalid sparse header")
    coordinates: set[tuple[int, int]] = set()
    row_indices: list[int] = []
    column_indices: list[int] = []
    values: list[Any] = []
    for row in rows[1:]:
        if len(row) != 3:
            raise ValueError(f"Excel sparse sheet {field!r} has an invalid row")
        try:
            row_index, column_index = int(row[0]), int(row[1])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Excel sparse sheet {field!r} has invalid coordinates") from exc
        coordinate = (row_index, column_index)
        if coordinate in coordinates or not (0 <= row_index < shape[0] and 0 <= column_index < shape[1]):
            raise ValueError(f"Excel sparse sheet {field!r} has duplicate or out-of-range coordinates")
        coordinates.add(coordinate)
        row_indices.append(row_index)
        column_indices.append(column_index)
        values.append(row[2])
    try:
        numeric = np.asarray(values, dtype=dtype)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Excel sparse sheet {field!r} contains non-numeric data") from exc
    return sparse.coo_matrix((numeric, (row_indices, column_indices)), shape=shape).tocsc()


def read_excel(source: str | PathLike[str]) -> MatpowerCase:
    """Read and validate a pwrs case workbook."""
    openpyxl = _openpyxl()
    path = _path(source)
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    if "_meta" not in workbook.sheetnames:
        raise ValueError("Excel case workbook is missing the '_meta' sheet")
    meta_rows = _rows(workbook["_meta"])
    if not meta_rows or meta_rows[0] != _META_HEADER:
        raise ValueError("Excel case workbook has an invalid '_meta' header")

    metadata: dict[str, tuple[Any, Any, Any, Any]] = {}
    for row in meta_rows[1:]:
        if len(row) != 5 or not isinstance(row[0], str) or row[0] in metadata:
            raise ValueError("Excel case workbook has invalid or duplicate metadata")
        metadata[row[0]] = (row[1], row[2], row[3], row[4])
    schema = metadata.pop("__schema__", None)
    version = metadata.pop("__schema_version__", None)
    if schema != ("schema", None, None, SCHEMA_NAME) or version != (
        "schema", None, None, SCHEMA_VERSION
    ):
        raise ValueError("unsupported Excel case schema")

    expected_sheets = {"_meta"}
    data: dict[str, Any] = {}
    for field, (kind, dtype_name, shape_text, value) in metadata.items():
        if kind == "scalar":
            data[field] = value
            continue
        expected_sheets.add(field)
        if field not in workbook.sheetnames:
            raise ValueError(f"Excel case workbook is missing sheet {field!r}")
        sheet = workbook[field]
        if kind == "dense":
            shape = _parse_shape(field, shape_text, 2)
            data[field] = _read_dense(
                field, sheet, _dtype(field, dtype_name), (shape[0], shape[1])
            )
        elif kind == "vector":
            shape = _parse_shape(field, shape_text, 1)
            data[field] = _read_vector(
                field, sheet, _dtype(field, dtype_name), (shape[0],)
            )
        elif kind == "sparse":
            shape = _parse_shape(field, shape_text, 2)
            data[field] = _read_sparse(
                field, sheet, _dtype(field, dtype_name), (shape[0], shape[1])
            )
        elif kind == "strings":
            shape = _parse_shape(field, shape_text, 1)
            rows = _rows(sheet)
            if not rows or rows[0] != ("value",) or len(rows) - 1 != shape[0]:
                raise ValueError(f"Excel sheet {field!r} does not match declared shape {shape}")
            data[field] = ["" if row[0] is None else str(row[0]) for row in rows[1:]]
        else:
            raise ValueError(f"Excel field {field!r} has unsupported kind {kind!r}")
    extra_sheets = sorted(set(workbook.sheetnames).difference(expected_sheets))
    if extra_sheets:
        raise ValueError(f"Excel case workbook has unregistered sheets: {', '.join(extra_sheets)}")
    return normalize_case(data)


__all__ = ["read_excel", "write_excel"]
